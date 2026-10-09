import csv
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import ALL_TEMP_FILES


ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("duco_tls_report", ROOT / "duco_tls_report.py")
report = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = report
spec.loader.exec_module(report)


class ReportTests(unittest.TestCase):
    def write_csv(self, path, rows, *, legacy=False):
        columns = [
            "timestamp", "phase", "protocol", "request", "http_status", "latency_ms",
            "error", "run_id", "repeat", "host", "endpoint",
        ]
        with path.open("w", newline="", encoding="utf-8") as target:
            writer = csv.writer(target)
            writer.writerow(columns[:7] if legacy else columns)
            for row in rows:
                writer.writerow(row[:7] if legacy else row)

    def rows(self):
        return [
            ["2026-10-08T09:05:00+00:00", "warm_sequential", "https", 1, 200, 10, "", "run-a", 1, "host-a", "info"],
            ["2026-10-08T09:10:00+00:00", "warm_sequential", "https", 2, 0, 8000, "timeout", "run-a", 1, "host-a", "info"],
            ["2026-10-08T10:05:00+00:00", "cold_handshakes", "https", 3, 200, 20, "", "run-a", 2, "host-a", "info"],
            ["2026-10-09T11:05:00+02:00", "parallel", "https", 4, 200, 30, "body read failed", "run-a", 2, "host-a", "info"],
        ]

    def diagnostic_row(self, **values):
        row = dict(zip([
            "timestamp", "phase", "protocol", "request", "http_status", "latency_ms",
            "error", "run_id", "repeat", "host", "endpoint",
        ], self.rows()[0]))
        row.update(record_type="request", scenario="parallel")
        row.update(values)
        return row

    def recovery_row(self, recovery_id, event, **values):
        return self.diagnostic_row(
            phase="recovery", request="", http_status="", latency_ms="", error="",
            record_type="recovery", recovery_id=recovery_id, recovery_event=event, **values,
        )

    def write_diagnostics(self, path, rows):
        columns = [
            "timestamp", "phase", "protocol", "request", "http_status", "latency_ms",
            "error", "run_id", "repeat", "host", "endpoint",
        ] + [
            "record_type", "scenario", "exception_type", "error_stage",
            "dns_ms", "connect_ms", "queue_ms", "http_ms", "body_ms",
            "new_connections", "established_connections", "reused_connections",
            "active_requests", "batch_id", "recovery_id", "recovery_event",
            "recovery_elapsed_s", "recovery_success_streak", "recovery_failure_streak",
            "first_failure_utc", "first_success_utc", "confirmed_utc", "recovery_outcome",
            "batch_size",
            "request_start_utc", "interval_s", "capture_manifest",
        ]
        with path.open("w", newline="", encoding="utf-8") as target:
            writer = csv.DictWriter(target, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)

    def sheet_records(self, workbook, title):
        for sheet in workbook:
            if sheet.title == title or sheet.title.startswith(title + "_"):
                rows = iter(sheet.values)
                headers = next(rows)
                for values in rows:
                    yield dict(zip(headers, values))

    def test_overview_aggregates_hours_days_repeats_and_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_csv(source, self.rows())
            original = source.read_bytes()
            self.assertEqual(report.export_report([source], output), 4)
            workbook = load_workbook(output)
            try:
                summary = dict(workbook["Overzicht"].values)
                self.assertEqual(summary["Requests"], 4)
                self.assertEqual(summary["Fouten"], 2)
                self.assertEqual(summary["Foutpercentage"], 0.5)
                self.assertEqual(summary["Gemiddelde ms (OK)"], 15)
                self.assertEqual(workbook["Fouten"].max_row, 3)
                hours = list(workbook["Uurtrend"].values)
                self.assertEqual(hours[1][1:5], (2, 1, 1, 0.5))
                self.assertEqual(hours[3][0], "2026-10-09T09:00:00.000Z")
                repeats = list(workbook["Herhalingen"].values)
                self.assertEqual(len(repeats), 3)
                self.assertEqual(repeats[1][4:7], (2, 1, 1))
                self.assertEqual(len(workbook["Uurtrend"]._charts), 2)
                self.assertEqual(workbook["Fouten"].freeze_panes, "A2")
                self.assertEqual(workbook["Fouten"].auto_filter.ref, "A1:AN3")
            finally:
                workbook.close()
            self.assertEqual(source.read_bytes(), original)

    def test_fault_sheets_roll_over_before_excel_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            row = self.rows()[1]
            self.write_csv(source, [row] * 5)
            report.export_report([source], output, max_rows=3)
            workbook = load_workbook(output)
            try:
                sheets = [sheet for sheet in workbook if sheet.title.startswith("Fouten")]
                self.assertEqual(len(sheets), 3)
                self.assertEqual([sheet.max_row for sheet in sheets], [3, 3, 2])
                self.assertEqual(sum(sheet.max_row - 1 for sheet in sheets), 5)
            finally:
                workbook.close()

    def test_old_repeat_files_can_be_combined(self):
        with tempfile.TemporaryDirectory() as directory:
            sources = [Path(directory) / f"repeat_{number}.csv" for number in (1, 2)]
            for source in sources:
                self.write_csv(source, self.rows()[:1], legacy=True)
            output = Path(directory) / "combined.xlsx"
            self.assertEqual(report.export_report(sources, output), 2)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["Herhalingen"].max_row, 3)
                self.assertEqual(dict(workbook["Overzicht"].values)["Bronbestanden"], 2)
            finally:
                workbook.close()

    def test_errors_are_text_not_excel_formulas(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            row = self.rows()[1]
            row[6] = '=HYPERLINK("https://example.invalid")'
            self.write_csv(source, [row])
            report.export_report([source], output)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["Fouten"]["K2"].data_type, "s")
                self.assertEqual(workbook["Fouten"]["K2"].value, row[6])
            finally:
                workbook.close()

    def test_save_failure_keeps_previous_report(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_csv(source, self.rows())
            output.write_bytes(b"previous report")
            with patch.object(report.os, "replace", side_effect=PermissionError("locked by Excel")):
                with self.assertRaises(PermissionError):
                    report.export_report([source], output)
            self.assertEqual(output.read_bytes(), b"previous report")
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [output])

    def test_empty_log_can_be_exported(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_csv(source, [])
            self.assertEqual(report.export_report([source], output), 0)
            workbook = load_workbook(output)
            try:
                self.assertEqual(dict(workbook["Overzicht"].values)["Requests"], 0)
                self.assertEqual(workbook["Fouten"].max_row, 1)
            finally:
                workbook.close()

    def test_invalid_data_does_not_replace_existing_report(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            row = self.rows()[0]
            row[5] = "NaN"
            self.write_csv(source, [row])
            output.write_bytes(b"previous report")
            temporary_files_before = set(ALL_TEMP_FILES)
            with self.assertRaisesRegex(ValueError, "Invalid latency"):
                report.export_report([source], output)
            self.assertEqual(output.read_bytes(), b"previous report")
            self.assertEqual(set(ALL_TEMP_FILES), temporary_files_before)

    def test_incomplete_last_physical_line_is_not_exported(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_csv(source, self.rows()[:1])
            with source.open("a", encoding="utf-8") as target:
                target.write("2026-10-08T09:15:00")
            self.assertEqual(report.export_report([source], output), 1)

    def test_snapshot_does_not_read_later_appends(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            self.write_csv(source, self.rows()[:1])
            lines = report.snapshot_lines(source)
            header = next(lines)
            with source.open("a", encoding="utf-8") as target:
                target.write("additional,data\n")
            self.assertEqual(len([header] + list(lines)), 2)

    def test_globs_with_no_matches_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "No CSV files"):
                report.expand_inputs([str(Path(directory) / "missing*.csv")])

    def test_diagnostics_connections_and_cancellations(self):
        rows = [
            self.diagnostic_row(new_connections=1, established_connections=1,
                                reused_connections=0, active_requests=1, batch_id=0),
            self.diagnostic_row(request=2, http_status=0, latency_ms=50, error="timeout",
                                exception_type="TimeoutError", error_stage="dns",
                                dns_ms=2.5, connect_ms=20, queue_ms=0, http_ms=3.5, body_ms=0,
                                new_connections=1, established_connections=0,
                                reused_connections=0, active_requests=4, batch_id=2),
            self.diagnostic_row(request=3, error="body read failed", error_stage="response_body",
                                exception_type="ClientPayloadError", http_ms=10, body_ms=4,
                                new_connections=0, established_connections=0,
                                reused_connections=1, active_requests=2),
            self.diagnostic_row(request=4, http_status=503, error="unavailable",
                                error_stage="http_status", new_connections=2,
                                established_connections=1, reused_connections=1, active_requests=3),
            self.diagnostic_row(request=5, new_connections=0, established_connections=0,
                                reused_connections=1, active_requests=0, batch_id=0),
            self.diagnostic_row(request=6, record_type="cancelled", exception_type="CancelledError",
                                error_stage="queue", new_connections=0, established_connections=0,
                                reused_connections=0, active_requests=3),
        ]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            self.assertEqual(report.export_report([source], output), 6)
            workbook = load_workbook(output)
            try:
                summary = dict(workbook["Overzicht"].values)
                self.assertEqual((summary["Requests"], summary["OK"], summary["Fouten"]), (6, 2, 4))
                scenarios = {row["Recordtype"]: row for row in self.sheet_records(workbook, "Scenarios")}
                requests = scenarios["request"]
                self.assertEqual((requests["Requests"], requests["OK"], requests["Fouten"]), (5, 2, 3))
                self.assertEqual(requests["Nieuwe connectiepogingen"], 4)
                self.assertEqual(requests["Opgezette connecties"], 2)
                self.assertEqual(requests["Hergebruik"], 3)
                self.assertEqual(requests["Metingen connectiepogingen"], 5)
                self.assertEqual(requests["Gem. DNS ms"], 2.5)
                self.assertEqual(requests["Metingen DNS"], 1)
                self.assertEqual(requests["Gem. HTTP ms"], 6.75)
                self.assertEqual(requests["Gem. actieve requests bij start"], 2)
                self.assertEqual(requests["Max. actieve requests bij start"], 4)
                self.assertEqual(scenarios["cancelled"]["Geannuleerd"], 1)
                self.assertEqual(scenarios["cancelled"]["Fouten"], 1)
                failures = list(self.sheet_records(workbook, "Fouten"))
                self.assertEqual([row["Connectietype"] for row in failures],
                                 ["new", "reused", "mixed", "unknown"])
                self.assertEqual(failures[0]["DNS ms"], 2.5)
                self.assertEqual(failures[0]["Connectie ms (DNS+TCP+TLS)"], 20)
                self.assertEqual(failures[0]["Wachtrij ms"], 0)
                self.assertEqual(failures[0]["HTTP ms"], 3.5)
                self.assertEqual(failures[0]["Body ms"], 0)
                self.assertEqual(failures[0]["Actieve requests bij start"], 4)
                self.assertEqual(failures[1]["HTTP-status"], 200)
                self.assertIsNone(failures[1]["DNS ms"])
                diagnostics = list(self.sheet_records(workbook, "Foutdiagnose"))
                self.assertEqual(sum(row["Fouten"] for row in diagnostics), 4)
                timeout = next(row for row in diagnostics if row["Exceptiontype"] == "TimeoutError")
                self.assertEqual(timeout["Foutfase"], "dns")
                self.assertEqual(timeout["Actieve requests bij start"], "4")
            finally:
                workbook.close()

    def test_scenario_keys_keep_health_and_recovery_under_owner(self):
        rows = [self.diagnostic_row()]
        for values in [
            {"run_id": "run-b"}, {"host": "host-b"}, {"endpoint": "status"},
            {"scenario": "cold_handshakes"}, {"phase": "health_https"},
            {"phase": "recovery_https"}, {"record_type": "cancelled"},
        ]:
            rows.append(self.diagnostic_row(**values))
        rows.append(self.diagnostic_row(repeat=2))
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            report.export_report([source], output)
            workbook = load_workbook(output)
            try:
                scenarios = list(self.sheet_records(workbook, "Scenarios"))
                self.assertEqual(len(scenarios), 8)
                self.assertEqual(sum(row["Requests"] for row in scenarios), 9)
                for measurement in ("health", "recovery"):
                    row = next(row for row in scenarios if row["Meettype"] == measurement)
                    self.assertEqual(row["Scenario"], "parallel")
                    self.assertEqual(row["Recordtype"], "request")
            finally:
                workbook.close()

    def test_legacy_diagnostics_are_unknown_not_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            sources = [Path(directory) / "legacy.csv", Path(directory) / "original.csv"]
            self.write_csv(sources[0], self.rows()[1:2], legacy=True)
            self.write_csv(sources[1], self.rows()[1:2])
            output = Path(directory) / "log.xlsx"
            self.assertEqual(report.export_report(sources, output), 2)
            workbook = load_workbook(output)
            try:
                for row in self.sheet_records(workbook, "Scenarios"):
                    self.assertEqual(row["Scenario"], "warm_sequential")
                    for header in ("Nieuwe connectiepogingen", "Opgezette connecties", "Hergebruik",
                                   "Gem. actieve requests bij start", "Max. actieve requests bij start"):
                        self.assertIsNone(row[header])
                    self.assertEqual(row["Metingen connectiepogingen"], 0)
                    self.assertEqual(row["Metingen actieve requests"], 0)
                for row in self.sheet_records(workbook, "Foutdiagnose"):
                    self.assertEqual(row["Connectietype"], "unknown")
                    self.assertEqual(row["Foutfase"], "onbekend")
                    self.assertEqual(row["Exceptiontype"], "onbekend")
                    self.assertEqual(row["Actieve requests bij start"], "onbekend")
                for row in self.sheet_records(workbook, "Fouten"):
                    self.assertIsNone(row["DNS ms"])
                    self.assertIsNone(row["Actieve requests bij start"])
            finally:
                workbook.close()

    def test_excel_is_formatted_and_scenarios_keep_protocol_and_batch_size(self):
        rows = [self.diagnostic_row(batch_size=2), self.diagnostic_row(batch_size=4),
                self.diagnostic_row(protocol="http", batch_size=2),
                self.recovery_row("id", "recovered", recovery_outcome="temporary")]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            report.export_report([source], output)
            workbook = load_workbook(output)
            try:
                scenarios = [row for row in self.sheet_records(workbook, "Scenarios") if row["Requests"]]
                self.assertEqual({(row["Protocol"], row["Requests in batch"]) for row in scenarios},
                                 {("https", "2"), ("https", "4"), ("http", "2")})
                sheet = workbook["Scenarios"]
                self.assertFalse(sheet.sheet_view.showGridLines)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertTrue(sheet["A1"].alignment.wrap_text)
                self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00256D5B")
                self.assertEqual(sheet["A2"].fill.fgColor.rgb, "00F1F5F4")
                self.assertEqual(sheet.row_dimensions[1].height, 44)
                self.assertEqual(workbook["Herstel"]["O2"].font.color.rgb, "00217346")
                self.assertEqual(workbook["Herstel"]["L2"].number_format, "0.00")
            finally:
                workbook.close()

    def test_recovery_events_are_streamed_and_do_not_count_as_requests(self):
        first_failure = "2026-10-08T09:00:00Z"
        first_success = "2026-10-08T09:00:14Z"
        confirmed = "2026-10-08T09:00:16Z"
        rows = [
            self.diagnostic_row(http_status=0, error="initial failure"),
            self.recovery_row("transient-a", "started", timestamp="2026-10-08T09:00:10Z",
                              first_failure_utc=first_failure, recovery_elapsed_s=10),
            self.diagnostic_row(phase="recovery_https", recovery_id="transient-a",
                                http_status=0, error="still down"),
            self.diagnostic_row(phase="recovery_https", recovery_id="transient-a"),
            self.recovery_row("transient-a", "first_success", timestamp=first_success,
                              first_success_utc=first_success, recovery_elapsed_s=14,
                              recovery_success_streak=1, recovery_failure_streak=0),
            self.recovery_row("transient-a", "recovered", timestamp=confirmed,
                              confirmed_utc=confirmed, recovery_elapsed_s=16,
                              recovery_success_streak=3, recovery_failure_streak=0,
                              recovery_outcome="transient"),
            self.recovery_row("timeout-a", "started", first_failure_utc=first_failure),
            self.diagnostic_row(phase="recovery", recovery_id="timeout-a",
                                record_type="cancelled", error="cancelled"),
            self.recovery_row("timeout-a", "timeout", recovery_elapsed_s=20,
                              recovery_failure_streak=4),
            self.recovery_row("interrupted-a", "interrupted", recovery_elapsed_s=30),
            self.recovery_row("pending-a", "started"),
            self.recovery_row("transient-b", "recovered", recovery_elapsed_s=24,
                              recovery_outcome="transient"),
            self.recovery_row("transient-b", "recovered", recovery_elapsed_s=24,
                              recovery_outcome="transient"),
            self.diagnostic_row(phase="recovery", recovery_id="transient-b", error="body failure"),
        ]
        event_rows = [row for row in rows if row["record_type"] == "recovery"]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            self.assertEqual(report.export_report([source], output), 5)
            workbook = load_workbook(output)
            try:
                summary = dict(workbook["Overzicht"].values)
                self.assertEqual((summary["Requests"], summary["OK"], summary["Fouten"]), (5, 1, 4))
                for title in ("Uurtrend", "Uren", "Dagen", "Herhalingen", "Scenarios"):
                    self.assertEqual(sum(row["Requests"] for row in self.sheet_records(workbook, title)), 5)
                self.assertEqual(workbook["Fouten"].max_row, 5)
                events = list(self.sheet_records(workbook, "Herstel"))
                self.assertEqual([(row["Recovery ID"], row["Herstelevent"]) for row in events],
                                 [(row["recovery_id"], row["recovery_event"]) for row in event_rows])
                self.assertEqual(events[2]["Bevestigd UTC"], "2026-10-08T09:00:16.000Z")
                self.assertEqual(events[2]["Successtreak"], 3)
                self.assertEqual(events[2]["Foutstreak"], 0)
                stats = {row["Hersteluitkomst"]: row for row in
                         self.sheet_records(workbook, "Scenarioherstel")}
                self.assertEqual(set(stats), {"transient", "timeout", "interrupted", "in_progress"})
                self.assertEqual(stats["transient"]["Herstelseries"], 2)
                self.assertEqual(stats["transient"]["Herstelrequests"], 3)
                self.assertEqual(stats["transient"]["Fouten"], 2)
                self.assertEqual(stats["transient"]["Gem. herstel s"], 20)
                self.assertEqual(stats["transient"]["Min. herstel s"], 16)
                self.assertEqual(stats["transient"]["Max. herstel s"], 24)
                self.assertEqual(stats["timeout"]["Gem. herstel s"], 20)
                self.assertEqual(stats["timeout"]["Geannuleerd"], 1)
                self.assertEqual(stats["interrupted"]["Gem. herstel s"], 30)
                self.assertEqual(stats["in_progress"]["Duurmetingen"], 0)
                self.assertIsNone(stats["in_progress"]["Gem. herstel s"])
                event_summary = next(row for row in self.sheet_records(workbook, "Scenarios")
                                     if row["Recordtype"] == "recovery")
                self.assertEqual(event_summary["Requests"], 0)
                self.assertEqual(event_summary["Herstelevents"], len(event_rows))
                self.assertIsNone(event_summary["Gem. ms"])
                self.assertEqual(stats["transient"]["Eerste fout UTC"], "2026-10-08T09:00:00.000Z")
            finally:
                workbook.close()

    def test_recovery_duration_fallback_and_ids_are_scoped_to_run(self):
        rows = [
            self.recovery_row("same-id", "recovered", timestamp="2026-10-08T11:00:10+02:00"),
            self.recovery_row("same-id", "started", timestamp="2026-10-08T09:00:00Z",
                              first_failure_utc="2026-10-08T11:00:00+02:00"),
            self.recovery_row("same-id", "timeout", run_id="run-b", recovery_elapsed_s=60),
        ]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            self.assertEqual(report.export_report([source], output), 0)
            workbook = load_workbook(output)
            try:
                stats = {row["Run ID"]: row for row in self.sheet_records(workbook, "Scenarioherstel")}
                self.assertEqual(stats["run-a"]["Hersteluitkomst"], "recovered")
                self.assertEqual(stats["run-a"]["Gem. herstel s"], 10)
                self.assertEqual(stats["run-a"]["Laatste event UTC"], "2026-10-08T09:00:10.000Z")
                self.assertEqual(stats["run-b"]["Hersteluitkomst"], "timeout")
                self.assertEqual(stats["run-b"]["Gem. herstel s"], 60)
                self.assertEqual(workbook["Fouten"].max_row, 1)
                self.assertEqual(workbook["Uurtrend"].max_row, 1)
            finally:
                workbook.close()

    def test_invalid_diagnostics_preserve_output_and_cleanup(self):
        invalid = [(column, value) for column in report.FLOAT_DIAGNOSTICS
                   for value in ("NaN", "inf", "-0.1")]
        invalid += [(column, value) for column in report.INTEGER_DIAGNOSTICS
                    for value in ("NaN", "inf", "-1", "1.5", str(2 ** 63))]
        invalid += [(column, "not-a-timestamp") for column in report.RECOVERY_TIMESTAMPS]
        original_temporary_directory = tempfile.TemporaryDirectory
        with original_temporary_directory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            output.write_bytes(b"previous report")
            for column, value in invalid:
                with self.subTest(column=column, value=value):
                    self.write_diagnostics(source, [self.diagnostic_row(),
                                                   self.diagnostic_row(**{column: value})])
                    paths = []

                    def tracked_directory(*args, **kwargs):
                        temporary = original_temporary_directory(*args, **kwargs)
                        paths.append(Path(temporary.name))
                        return temporary

                    temporary_files_before = set(ALL_TEMP_FILES)
                    with patch.object(report.tempfile, "TemporaryDirectory", side_effect=tracked_directory):
                        with self.assertRaisesRegex(ValueError, r"CSV line 3"):
                            report.export_report([source], output)
                    self.assertEqual(output.read_bytes(), b"previous report")
                    self.assertEqual(list(Path(directory).glob("*.xlsx")), [output])
                    self.assertEqual(set(ALL_TEMP_FILES), temporary_files_before)
                    self.assertTrue(paths)
                    self.assertTrue(all(not path.exists() for path in paths))

    def test_recovery_metadata_is_validated_too(self):
        cases = [
            self.recovery_row("id", "timeout", recovery_elapsed_s="NaN"),
            self.recovery_row("id", "timeout", active_requests=-1),
            self.recovery_row("id", "timeout", recovery_failure_streak="1.5"),
            self.recovery_row("id", "timeout", confirmed_utc="invalid"),
            self.recovery_row("", "started"), self.recovery_row("id", "invalid-event"),
            self.diagnostic_row(record_type="invalid-type"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            output.write_bytes(b"previous report")
            for row in cases:
                with self.subTest(row=row):
                    self.write_diagnostics(source, [row])
                    before = set(ALL_TEMP_FILES)
                    with self.assertRaises(ValueError):
                        report.export_report([source], output)
                    self.assertEqual(output.read_bytes(), b"previous report")
                    self.assertEqual(set(ALL_TEMP_FILES), before)

    def test_workbook_save_failure_cleans_database_and_xml(self):
        original_temporary_directory = tempfile.TemporaryDirectory
        with original_temporary_directory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, [self.diagnostic_row(), self.recovery_row("id", "started")])
            output.write_bytes(b"previous report")
            paths = []

            def tracked_directory(*args, **kwargs):
                temporary = original_temporary_directory(*args, **kwargs)
                paths.append(Path(temporary.name))
                return temporary

            before = set(ALL_TEMP_FILES)
            with patch.object(report.tempfile, "TemporaryDirectory", side_effect=tracked_directory):
                with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("save failed")):
                    with self.assertRaisesRegex(OSError, "save failed"):
                        report.export_report([source], output)
            self.assertEqual(output.read_bytes(), b"previous report")
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [output])
            self.assertEqual(set(ALL_TEMP_FILES), before)
            self.assertTrue(all(not path.exists() for path in paths))

    def test_all_report_sheets_roll_over(self):
        rows = []
        for number in range(7):
            rows.extend([
                self.diagnostic_row(timestamp=f"2026-10-{number + 1:02d}T09:00:00Z",
                                    run_id=f"run-{number}", phase=f"stress_{number}",
                                    scenario=f"scenario-{number}", http_status=0, error="timeout"),
                self.recovery_row(f"id-{number}", "timeout", run_id=f"run-{number}",
                                  scenario=f"scenario-{number}", recovery_elapsed_s=number),
            ])
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            report.export_report([source], output, max_rows=3)
            workbook = load_workbook(output)
            try:
                for sheet in workbook:
                    self.assertLessEqual(sheet.max_row, 3)
                    self.assertEqual(sheet.auto_filter.ref,
                                     f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}")
                for title in ("Fouten", "Uurtrend", "Uren", "Dagen", "Herhalingen",
                              "Herstel", "Foutdiagnose", "Scenarioherstel"):
                    self.assertEqual(len(list(self.sheet_records(workbook, title))), 7)
                    self.assertIn(title + "_4", workbook.sheetnames)
                self.assertEqual(len(list(self.sheet_records(workbook, "Scenarios"))), 14)
                self.assertEqual(len(workbook["Uurtrend"]._charts), 0)
            finally:
                workbook.close()

    def test_aggregations_are_disk_backed_with_fixed_cache_and_streamed_rows(self):
        aggregates = report.DiskAggregates()
        directory = Path(aggregates.directory.name)
        connection = aggregates.connection
        try:
            database = Path(connection.execute("PRAGMA database_list").fetchone()[2])
            self.assertEqual(database.parent, directory)
            self.assertTrue(database.is_file())
            self.assertEqual(connection.execute("PRAGMA cache_size").fetchone()[0], -2048)
            self.assertEqual(connection.execute("PRAGMA temp_store").fetchone()[0], 1)
            self.assertEqual(connection.execute("PRAGMA temp.cache_size").fetchone()[0], -2048)
            self.assertEqual(connection.execute("PRAGMA mmap_size").fetchone()[0], 0)
            moment = report.utc_moment("2026-10-08T09:00:00Z")
            recovery = report.parse_record(self.recovery_row("id", "timeout", recovery_elapsed_s=5))
            dictionary_sizes = {name: len(value) for name, value in vars(aggregates).items()
                                if isinstance(value, dict)}
            for number in range(1500):
                run_id = f"run-{number:04d}"
                aggregates.add("repeats", (run_id, 1, "host", "info"), moment, 1, True)
                aggregates.add_recovery((run_id, "host", "info", "parallel", str(number)),
                                        recovery, moment, False)
            rows = aggregates.rows("repeats")
            self.assertIs(iter(rows), rows)
            self.assertEqual(sum(1 for row in rows), 1500)
            self.assertEqual(sum(1 for row in aggregates.recovery_rows()), 1500)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM recovery_episodes").fetchone()[0], 1500)
            self.assertEqual(dictionary_sizes, {name: len(value) for name, value in vars(aggregates).items()
                                               if isinstance(value, dict)})
        finally:
            aggregates.close()
        self.assertFalse(directory.exists())

    def test_partial_diagnostic_sums_include_coverage_not_assumed_zeroes(self):
        rows = [
            self.diagnostic_row(new_connections=2, active_requests=4, error="body error"),
            self.diagnostic_row(established_connections=1, active_requests=2),
            self.diagnostic_row(),
        ]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "log.csv"
            output = Path(directory) / "log.xlsx"
            self.write_diagnostics(source, rows)
            report.export_report([source], output)
            workbook = load_workbook(output)
            try:
                row = next(self.sheet_records(workbook, "Scenarios"))
                self.assertEqual(row["Requests"], 3)
                self.assertEqual(row["Nieuwe connectiepogingen"], 2)
                self.assertEqual(row["Opgezette connecties"], 1)
                self.assertIsNone(row["Hergebruik"])
                self.assertEqual(row["Metingen connectiepogingen"], 1)
                self.assertEqual(row["Metingen opgezette connecties"], 1)
                self.assertEqual(row["Metingen hergebruik"], 0)
                self.assertEqual(row["Gem. actieve requests bij start"], 3)
                self.assertEqual(row["Max. actieve requests bij start"], 4)
                self.assertEqual(row["Metingen actieve requests"], 2)
                self.assertEqual(next(self.sheet_records(workbook, "Fouten"))["Connectietype"], "unknown")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()