import asyncio
import csv
import errno
import io
import socket
import ssl
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import aiohttp
from aiohttp import web
from openpyxl import load_workbook

import duco_tls_stress as stress


def arguments(output, **overrides):
    values = dict(output=str(output), host="example.invalid", endpoint="info",
                  verbose=False, duration=30, concurrency=4, interval=1,
                  timeout=8, max_failures=2, repeat=1, repeat_delay=0,
                  excel_output=None, excel_interval=3600, no_excel=True,
                  scenario="all", recovery_duration=95, resume_after_recovery=False)
    values.update(overrides)
    return SimpleNamespace(**values)


class Clock:
    now = 0.0

    def monotonic(self):
        return self.now

    async def sleep(self, seconds):
        self.now += seconds

    def timestamp(self):
        return (datetime(2026, 10, 8, tzinfo=timezone.utc) + timedelta(seconds=self.now)).isoformat()


class Session:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        return False


class DiagnosticsTests(unittest.TestCase):
    def read_rows(self, output):
        with output.open(newline="", encoding="utf-8") as source:
            return list(csv.DictReader(source))

    def test_recovery_requires_three_successes_and_resets_after_relapse(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            test = stress.DucoStressTest(arguments(output))
            clock = Clock()
            outcomes = iter([False, True, True, False, True, True, True])
            probes = []

            async def request(session, phase, protocol):
                probes.append((clock.now, protocol))
                return next(outcomes) if protocol == "https" else True

            test.first_failure = dict(timestamp=clock.timestamp(), monotonic=0, scenario="parallel")
            with patch.object(stress.time, "monotonic", clock.monotonic), \
                 patch.object(stress, "timestamp", clock.timestamp), \
                 patch.object(stress.asyncio, "sleep", clock.sleep), \
                 patch.object(test, "request", request), \
                 patch.object(test, "session", return_value=Session()) as sessions, \
                 redirect_stdout(io.StringIO()):
                self.assertTrue(asyncio.run(test.recovery()))
            test.close()
            rows = self.read_rows(output)
            checks = [row for row in rows if row["recovery_event"] == "probe"]
            self.assertEqual([int(row["recovery_success_streak"]) for row in checks], [0, 1, 2, 0, 1, 2, 3])
            self.assertEqual([int(row["recovery_failure_streak"]) for row in checks], [2, 0, 0, 1, 0, 0, 0])
            self.assertEqual([moment for moment, protocol in probes if protocol == "https"], [10, 20, 30, 40, 50, 60, 70])
            self.assertEqual(len(probes), 14)
            self.assertTrue(all(call.kwargs["force_close"] for call in sessions.call_args_list))
            self.assertEqual(rows[-1]["recovery_outcome"], "temporary")
            self.assertEqual(float(rows[-1]["recovery_elapsed_s"]), 70)
            self.assertEqual(rows[-1]["first_success_utc"], "2026-10-08T00:00:20+00:00")
            self.assertEqual(rows[-1]["confirmed_utc"], "2026-10-08T00:01:10+00:00")
            self.assertEqual(sum(row["recovery_event"] == "first_success" for row in rows), 1)
            self.assertEqual(len({row["recovery_id"] for row in rows}), 1)
            stress.export_report([output], output.with_suffix(".xlsx"))
            workbook = load_workbook(output.with_suffix(".xlsx"))
            try:
                summary = dict(workbook["Overzicht"].values)
                self.assertEqual(summary["Requests"], 0)
                self.assertEqual(workbook["Herstel"].max_row, len(rows) + 1)
                statistics = list(workbook["Scenarioherstel"].values)
                result = dict(zip(statistics[0], statistics[1]))
                self.assertEqual(result["Hersteluitkomst"], "temporary")
                self.assertEqual(result["Gem. herstel s"], 70)
            finally:
                workbook.close()

    def test_timeout_and_deadline_cap_each_probe(self):
        for duration, expected_probes, request_time in [(35, 6, 0), (12, 1, 8)]:
            with self.subTest(duration=duration), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "log.csv"
                test = stress.DucoStressTest(arguments(output, recovery_duration=duration))
                clock = Clock()
                session_options = []
                probes = []

                def session(**kwargs):
                    session_options.append(kwargs)
                    return Session()

                async def request(session, phase, protocol):
                    probes.append(protocol)
                    clock.now += min(request_time, session_options[-1]["timeout"])
                    return False

                test.first_failure = dict(timestamp=clock.timestamp(), monotonic=0, scenario="cold_handshakes")
                with patch.object(stress.time, "monotonic", clock.monotonic), \
                     patch.object(stress.asyncio, "sleep", clock.sleep), \
                     patch.object(test, "session", session), \
                     patch.object(test, "request", request), redirect_stdout(io.StringIO()):
                    self.assertFalse(asyncio.run(test.recovery()))
                test.close()
                rows = self.read_rows(output)
                self.assertEqual(len(probes), expected_probes)
                self.assertEqual(clock.now, duration)
                self.assertEqual(rows[-1]["recovery_outcome"], "not_recovered")
                self.assertEqual(float(rows[-1]["recovery_elapsed_s"]), duration)
                self.assertTrue(all(options["timeout"] <= 8 for options in session_options))

    def test_interrupted_recovery_records_outcome(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            test = stress.DucoStressTest(arguments(output))
            test.first_failure = dict(timestamp=stress.timestamp(), monotonic=stress.time.monotonic(), scenario="parallel")
            with patch.object(stress.asyncio, "sleep", AsyncMock(side_effect=asyncio.CancelledError())), \
                 redirect_stdout(io.StringIO()), self.assertRaises(asyncio.CancelledError):
                asyncio.run(test.recovery())
            test.close()
            self.assertEqual(self.read_rows(output)[-1]["recovery_outcome"], "interrupted")

    def test_recovery_never_resumes_without_explicit_option(self):
        for resume, repeats, recovered, expected_cycles, expected_exit in [
            (False, 2, True, 1, 0), (True, 2, True, 2, 0),
            (True, 1, True, 1, 0), (True, 2, False, 1, 1),
        ]:
            with self.subTest(resume=resume, repeats=repeats, recovered=recovered), tempfile.TemporaryDirectory() as directory:
                calls = []

                async def phase(test, seconds):
                    calls.append(seconds)
                    if len(calls) == 1:
                        test.first_failure = dict(timestamp=stress.timestamp(), monotonic=stress.time.monotonic(), scenario="warm_sequential")
                        raise RuntimeError("failure")

                with patch.object(stress.DucoStressTest, "phase_warm", phase), \
                     patch.object(stress.DucoStressTest, "health", AsyncMock()), \
                     patch.object(stress.DucoStressTest, "recovery", AsyncMock(return_value=recovered)) as recovery, \
                     patch.object(stress.DucoStressTest, "phase_cold", AsyncMock()) as cold, \
                     redirect_stdout(io.StringIO()):
                    result = asyncio.run(stress.main(arguments(Path(directory) / "log.csv", scenario="warm-sequential", repeat=repeats,
                                                              resume_after_recovery=resume)))
                self.assertEqual(result, expected_exit)
                self.assertEqual(calls, [30] * expected_cycles)
                recovery.assert_awaited_once()
                cold.assert_not_awaited()

    def test_initial_https_health_failure_starts_recovery_without_stress(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            protocols = []

            async def request(test, session, phase, protocol):
                protocols.append(protocol)
                if protocol == "https":
                    test.first_failure = dict(timestamp=stress.timestamp(), monotonic=stress.time.monotonic(), scenario=test.scenario)
                    return False
                return True

            with patch.object(stress.DucoStressTest, "request", request), \
                 patch.object(stress.DucoStressTest, "session", return_value=Session()), \
                 patch.object(stress.DucoStressTest, "phase_warm", AsyncMock()) as warm, \
                 patch.object(stress.DucoStressTest, "recovery", AsyncMock(return_value=True)) as recovery, \
                 redirect_stdout(io.StringIO()):
                self.assertEqual(asyncio.run(stress.main(arguments(output))), 0)
            warm.assert_not_awaited()
            recovery.assert_awaited_once()
            self.assertEqual(protocols, ["https", "http"])

    def test_resume_all_starts_a_new_cycle_not_the_failed_phase(self):
        with tempfile.TemporaryDirectory() as directory:
            calls = []

            async def warm(test, seconds):
                calls.append("warm")

            async def cold(test, seconds):
                calls.append("cold")
                if calls == ["warm", "cold"]:
                    test.first_failure = dict(timestamp=stress.timestamp(), monotonic=stress.time.monotonic(), scenario=test.scenario)
                    raise RuntimeError("cold failure")

            async def parallel(test, seconds):
                calls.append("parallel")

            with patch.object(stress.DucoStressTest, "phase_warm", warm), \
                 patch.object(stress.DucoStressTest, "phase_cold", cold), \
                 patch.object(stress.DucoStressTest, "phase_parallel", parallel), \
                 patch.object(stress.DucoStressTest, "health", AsyncMock()), \
                 patch.object(stress.DucoStressTest, "recovery", AsyncMock(return_value=True)), \
                 redirect_stdout(io.StringIO()):
                asyncio.run(stress.main(arguments(Path(directory) / "log.csv", repeat=2, resume_after_recovery=True)))
            self.assertEqual(calls, ["warm", "cold", "warm", "cold", "parallel"])

    def test_each_scenario_runs_independently(self):
        for scenario, selected in [("warm-sequential", "phase_warm"), ("cold-handshakes", "phase_cold"), ("parallel", "phase_parallel")]:
            with self.subTest(scenario=scenario), tempfile.TemporaryDirectory() as directory:
                with patch.object(stress.DucoStressTest, "phase_warm", AsyncMock()) as warm, \
                     patch.object(stress.DucoStressTest, "phase_cold", AsyncMock()) as cold, \
                     patch.object(stress.DucoStressTest, "phase_parallel", AsyncMock()) as parallel, \
                     patch.object(stress.DucoStressTest, "health", AsyncMock()) as health, redirect_stdout(io.StringIO()):
                    asyncio.run(stress.main(arguments(Path(directory) / "log.csv", scenario=scenario)))
                for name, phase in [("phase_warm", warm), ("phase_cold", cold), ("phase_parallel", parallel)]:
                    if name == selected:
                        phase.assert_awaited_once_with(30)
                    else:
                        phase.assert_not_awaited()
                self.assertEqual(health.await_count, 2)

    def test_parallel_failure_cancels_pending_requests(self):
        with tempfile.TemporaryDirectory() as directory:
            test = stress.DucoStressTest(arguments(Path(directory) / "log.csv"))
            calls = []
            cancelled = []

            async def request(session, phase):
                number = len(calls)
                calls.append(number)
                if number == 0:
                    await asyncio.sleep(0)
                    return False
                try:
                    await asyncio.Event().wait()
                except asyncio.CancelledError:
                    cancelled.append(number)
                    raise

            with patch.object(test, "request", request), self.assertRaisesRegex(RuntimeError, "cancelling"):
                asyncio.run(test.parallel_batch(None))
            test.close()
            self.assertEqual(len(calls), 4)
            self.assertEqual(len(cancelled), 3)

    def test_exception_stage_classification_does_not_guess_tls(self):
        key = SimpleNamespace(host="example.invalid", port=443, ssl=False)
        exceptions = [
            (aiohttp.ClientConnectorDNSError(key, socket.gaierror("DNS failed")), "dns"),
            (aiohttp.ClientConnectorError(key, ConnectionRefusedError("refused")), "tcp_connect"),
            (aiohttp.ClientSSLError(key, ssl.SSLError("handshake")), "tls_handshake"),
            (aiohttp.ServerDisconnectedError(), "http_request"),
        ]
        for exception, stage in exceptions:
            with self.subTest(exception=type(exception).__name__):
                self.assertEqual(stress.error_stage(exception, stress.RequestTrace()), stage)
        self.assertEqual(stress.error_stage(TimeoutError(), stress.RequestTrace(stage="tcp_tls_connect")), "tcp_tls_connect")
        reset = aiohttp.ClientConnectorError(key, ConnectionResetError(errno.ECONNRESET, "reset"))
        self.assertEqual(stress.error_stage(reset, stress.RequestTrace(stage="tcp_tls_connect")), "tcp_tls_connect")
        self.assertEqual(stress.error_stage(OSError(), stress.RequestTrace(stage="response_body")), "response_body")

    def test_cold_requests_use_fresh_connections_and_five_or_ten_second_pause(self):
        for interval in (5, 10):
            with self.subTest(interval=interval), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "log.csv"
                test = stress.DucoStressTest(arguments(output, concurrency=1, interval=interval))
                clock = Clock()
                starts = []

                async def request(session, phase):
                    starts.append(clock.now)
                    clock.now += 1
                    return True

                with patch.object(stress.time, "monotonic", clock.monotonic), \
                     patch.object(stress.asyncio, "sleep", clock.sleep), \
                     patch.object(test, "request", request), \
                     patch.object(test, "session", side_effect=lambda **kwargs: Session()) as sessions, \
                     redirect_stdout(io.StringIO()):
                    asyncio.run(test.phase_cold(30))
                test.close()
                self.assertEqual(starts, list(range(0, 30, interval + 1)))
                self.assertEqual(sessions.call_count, len(starts))
                self.assertTrue(all(call.kwargs["force_close"] for call in sessions.call_args_list))

    def test_cancelled_requests_are_logged_without_replacing_first_failure(self):
        class PendingResponse:
            def __init__(self, number):
                self.number = number

            async def __aenter__(self):
                if self.number == 0:
                    await asyncio.sleep(0)
                    raise aiohttp.ServerDisconnectedError()
                await asyncio.Event().wait()

            async def __aexit__(self, *args):
                return False

        class PendingSession:
            calls = 0

            def get(self, *args, **kwargs):
                response = PendingResponse(self.calls)
                self.calls += 1
                return response

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            test = stress.DucoStressTest(arguments(output))
            test.scenario = "parallel"
            with redirect_stdout(io.StringIO()), self.assertRaises(RuntimeError):
                asyncio.run(test.parallel_batch(PendingSession()))
            test.close()
            rows = self.read_rows(output)
            self.assertEqual([row["record_type"] for row in rows].count("cancelled"), 3)
            self.assertEqual(rows[0]["exception_type"], "ServerDisconnectedError")
            self.assertEqual(test.first_failure["timestamp"], rows[0]["timestamp"])
            self.assertEqual(test.active_requests, 0)
            self.assertEqual({row["recovery_id"] for row in rows}, {test.recovery_id})
            self.assertEqual({row["batch_id"] for row in rows}, {"1"})
            self.assertEqual({row["active_requests"] for row in rows}, {"1", "2", "3", "4"})

    def test_real_tls_and_response_body_failures_are_distinguished(self):
        async def exercise(output):
            application = web.Application()

            async def truncated(request):
                response = web.StreamResponse(headers={"Content-Length": "100"})
                await response.prepare(request)
                await response.write(b"short")
                request.transport.close()
                return response

            application.router.add_get("/info", truncated)
            runner = web.AppRunner(application)
            await runner.setup()
            site = web.TCPSite(runner, "127.0.0.1", 0)
            await site.start()
            port = site._server.sockets[0].getsockname()[1]
            test = stress.DucoStressTest(arguments(output, host=f"127.0.0.1:{port}"))
            try:
                async with test.session(force_close=True) as session:
                    self.assertFalse(await test.request(session, "cold_handshakes", "https"))
                async with test.session(force_close=True) as session:
                    self.assertFalse(await test.request(session, "warm_sequential", "http"))
                await runner.cleanup()
                async with test.session(force_close=True) as session:
                    self.assertFalse(await test.request(session, "cold_handshakes", "http"))
            finally:
                test.close()
                await runner.cleanup()

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            with redirect_stdout(io.StringIO()):
                asyncio.run(exercise(output))
            rows = self.read_rows(output)
            self.assertEqual([row["error_stage"] for row in rows], ["tls_handshake", "response_body", "tcp_connect"])
            self.assertEqual(rows[1]["http_status"], "200")
            self.assertEqual(rows[1]["exception_type"], "ClientPayloadError")
            self.assertEqual([row["established_connections"] for row in rows], ["0", "1", "0"])
            self.assertTrue(all(row["connect_ms"] for row in rows))

    def test_actual_trace_reports_new_connections_reuse_and_response_errors(self):
        async def exercise(output):
            application = web.Application()

            async def info(request):
                return web.json_response({"ok": True})

            application.router.add_get("/info", info)
            runner = web.AppRunner(application)
            await runner.setup()
            site = web.TCPSite(runner, "127.0.0.1", 0)
            await site.start()
            port = site._server.sockets[0].getsockname()[1]
            test = stress.DucoStressTest(arguments(output, host=f"localhost:{port}"))
            try:
                async with test.session() as session:
                    self.assertTrue(await test.request(session, "warm_sequential", "http"))
                    self.assertTrue(await test.request(session, "warm_sequential", "http"))
                async with test.session(force_close=True) as session:
                    self.assertTrue(await test.request(session, "cold_handshakes", "http"))
            finally:
                test.close()
                await runner.cleanup()

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            asyncio.run(exercise(output))
            rows = self.read_rows(output)
            self.assertEqual([row["new_connections"] for row in rows], ["1", "0", "1"])
            self.assertEqual([row["established_connections"] for row in rows], ["1", "0", "1"])
            self.assertEqual([row["reused_connections"] for row in rows], ["0", "1", "0"])
            self.assertNotEqual(rows[0]["dns_ms"], "")
            self.assertEqual(rows[1]["dns_ms"], "")
            self.assertTrue(all(float(row["http_ms"]) >= 0 for row in rows))
            self.assertTrue(all(float(row["body_ms"]) >= 0 for row in rows))
            self.assertTrue(all(None not in row for row in rows))

    def test_aiohttp_does_not_silently_retry_disconnects(self):
        async def exercise(output):
            calls = []
            application = web.Application()

            async def disconnect(request):
                calls.append(request)
                if len(calls) == 1:
                    return web.json_response({"ok": True})
                request.transport.close()
                return web.Response()

            application.router.add_get("/info", disconnect)
            runner = web.AppRunner(application)
            await runner.setup()
            site = web.TCPSite(runner, "127.0.0.1", 0)
            await site.start()
            port = site._server.sockets[0].getsockname()[1]
            test = stress.DucoStressTest(arguments(output, host=f"127.0.0.1:{port}"))
            try:
                async with test.session() as session:
                    self.assertTrue(await test.request(session, "warm_sequential", "http"))
                    self.assertFalse(await test.request(session, "warm_sequential", "http"))
            finally:
                test.close()
                await runner.cleanup()
            self.assertEqual(len(calls), 2)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "log.csv"
            with redirect_stdout(io.StringIO()):
                asyncio.run(exercise(output))
            failure = self.read_rows(output)[-1]
            self.assertEqual(failure["exception_type"], "ServerDisconnectedError")
            self.assertEqual(failure["error_stage"], "http_request")
            self.assertEqual(failure["reused_connections"], "1")
            self.assertEqual(failure["new_connections"], "0")


if __name__ == "__main__":
    unittest.main()