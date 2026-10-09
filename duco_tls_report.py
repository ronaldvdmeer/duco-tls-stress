#!/usr/bin/env python3
"""Stream Duco CSV logs into one Excel overview without loading raw data in memory."""

import argparse
import csv
import glob
import math
import os
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path


DEFAULT_INPUT = Path(__file__).resolve().parent / "logs" / "duco_tls.csv"
EXCEL_MAX_ROWS = 1_048_576
REQUIRED_COLUMNS = {
    "timestamp", "phase", "protocol", "request",
    "http_status", "latency_ms", "error",
}
DIAGNOSTIC_COLUMNS = [
    "record_type", "scenario", "exception_type", "error_stage",
    "dns_ms", "connect_ms", "queue_ms", "http_ms", "body_ms",
    "new_connections", "established_connections", "reused_connections",
    "active_requests", "batch_id", "recovery_id", "recovery_event",
    "recovery_elapsed_s", "recovery_success_streak", "recovery_failure_streak",
    "first_failure_utc", "first_success_utc", "confirmed_utc", "recovery_outcome",
    "batch_size",
    "request_start_utc", "interval_s", "capture_manifest",
]
DIAGNOSTIC_HEADERS = [
    "Recordtype", "Scenario", "Exceptiontype", "Foutfase",
    "DNS ms", "Connectie ms (DNS+TCP+TLS)", "Wachtrij ms", "HTTP ms", "Body ms",
    "Nieuwe connectiepogingen", "Opgezette connecties", "Hergebruik",
    "Actieve requests bij start", "Batch ID", "Recovery ID", "Herstelevent",
    "Herstel verstreken s", "Successtreak", "Foutstreak",
    "Eerste fout UTC", "Eerste succes UTC", "Bevestigd UTC", "Hersteluitkomst",
    "Requests in batch",
    "Requeststart UTC", "Interval s", "Capture metadata",
]
FLOAT_DIAGNOSTICS = {
    "dns_ms", "connect_ms", "queue_ms", "http_ms", "body_ms", "recovery_elapsed_s",
    "interval_s",
}
INTEGER_DIAGNOSTICS = {
    "new_connections", "established_connections", "reused_connections",
    "active_requests", "batch_id", "recovery_success_streak", "recovery_failure_streak",
    "batch_size",
}
RECOVERY_TIMESTAMPS = {"first_failure_utc", "first_success_utc", "confirmed_utc", "request_start_utc"}
RECOVERY_EVENTS = {"started", "probe", "first_success", "recovered", "timeout", "interrupted"}
SCENARIO_KEYS = [("run_id", "TEXT"), ("host", "TEXT"), ("endpoint", "TEXT"),
                 ("scenario", "TEXT"), ("measurement", "TEXT"), ("record_type", "TEXT"),
                 ("protocol", "TEXT"), ("batch_size", "TEXT"), ("interval_s", "TEXT"),
                 ("capture_manifest", "TEXT")]
SCENARIO_HEADERS = ["Run ID", "Host", "Endpoint", "Scenario", "Meettype", "Recordtype",
                    "Protocol", "Requests in batch", "Interval s", "Capture metadata"]
ERROR_COLUMNS = [
    "Tijd UTC", "Run ID", "Repeat", "Host", "Endpoint", "Fase",
    "Protocol", "Request", "HTTP-status", "Latency ms", "Fout", "Bronbestand",
] + DIAGNOSTIC_HEADERS + ["Connectietype"]
RECOVERY_COLUMNS = [
    "Tijd UTC", "Run ID", "Repeat", "Host", "Endpoint", "Recovery ID", "Scenario",
    "Herstelevent", "Eerste fout UTC", "Eerste succes UTC", "Bevestigd UTC",
    "Herstel verstreken s", "Successtreak", "Foutstreak", "Hersteluitkomst", "Bronbestand",
]
SCENARIO_EXTRA_HEADERS = [
    "Nieuwe connectiepogingen", "Opgezette connecties", "Hergebruik",
    "Metingen connectiepogingen", "Metingen opgezette connecties", "Metingen hergebruik",
    "Geannuleerd", "Herstelevents", "Gem. actieve requests bij start",
    "Max. actieve requests bij start", "Metingen actieve requests",
    "Gem. DNS ms", "Metingen DNS", "Gem. connectie ms (DNS+TCP+TLS)", "Metingen connectie",
    "Gem. wachtrij ms", "Metingen wachtrij", "Gem. HTTP ms", "Metingen HTTP",
    "Gem. body ms", "Metingen body",
]
STAGE_COLUMNS = ["dns_ms", "connect_ms", "queue_ms", "http_ms", "body_ms"]
METRIC_COLUMNS = [
    "Requests", "OK", "Fouten", "Foutpercentage", "Gem. ms", "Max. ms",
    "Gem. OK ms", "Eerste meting UTC", "Laatste meting UTC",
]


@dataclass
class Metrics:
    requests: int = 0
    successes: int = 0
    total_ms: float = 0.0
    max_ms: float = 0.0
    success_ms: float = 0.0
    first: datetime = None
    last: datetime = None

    def add(self, moment, latency, success):
        self.requests += 1
        self.successes += int(success)
        self.total_ms += latency
        self.max_ms = max(self.max_ms, latency)
        if success:
            self.success_ms += latency
        self.first = min(self.first, moment) if self.first else moment
        self.last = max(self.last, moment) if self.last else moment

    def values(self):
        failures = self.requests - self.successes
        return [
            self.requests, self.successes, failures,
            failures / self.requests if self.requests else 0,
            round(self.total_ms / self.requests, 2) if self.requests else 0,
            round(self.max_ms, 2),
            round(self.success_ms / self.successes, 2) if self.successes else None,
            self.first, self.last,
        ]


class DiskAggregates:
    CACHE_KIB = 2048
    KEYS = {
        "hours": [("hour", "TEXT")],
        "hourly_phases": [("hour", "TEXT"), ("host", "TEXT"), ("endpoint", "TEXT"),
                          ("phase", "TEXT"), ("protocol", "TEXT")],
        "daily_phases": [("day", "TEXT"), ("host", "TEXT"), ("endpoint", "TEXT"),
                         ("phase", "TEXT"), ("protocol", "TEXT")],
        "repeats": [("run_id", "TEXT"), ("repeat", "INTEGER"), ("host", "TEXT"),
                    ("endpoint", "TEXT")],
        "scenarios": SCENARIO_KEYS,
        "diagnostics": SCENARIO_KEYS + [
            ("error_stage", "TEXT"), ("exception_type", "TEXT"),
            ("connection_type", "TEXT"), ("active_requests", "TEXT"),
        ],
    }
    SCENARIO_REDUCERS = {
        "new_connections": "SUM", "established_connections": "SUM",
        "reused_connections": "SUM", "new_samples": "SUM", "established_samples": "SUM",
        "reused_samples": "SUM", "cancelled": "SUM", "events": "SUM",
        "active_total": "SUM", "active_max": "MAX", "active_samples": "SUM",
        **{name: "SUM" for stage in STAGE_COLUMNS for name in (stage + "_total", stage + "_samples")},
    }
    RECOVERY_KEYS = SCENARIO_KEYS[:4] + [("recovery_id", "TEXT")]

    def __init__(self):
        self.directory = tempfile.TemporaryDirectory(prefix="duco_report_")
        self.connection = None
        self.pending = 0
        self.upserts = {}
        try:
            self.connection = sqlite3.connect(Path(self.directory.name) / "aggregates.sqlite3")
            self.connection.execute(f"PRAGMA cache_size = -{self.CACHE_KIB}")
            self.connection.execute("PRAGMA temp_store = FILE")
            self.connection.execute(f"PRAGMA temp.cache_size = -{self.CACHE_KIB}")
            self.connection.execute("PRAGMA mmap_size = 0")
            for table, keys in self.KEYS.items():
                names = ", ".join(name for name, kind in keys)
                definitions = ", ".join(f"{name} {kind} NOT NULL" for name, kind in keys)
                reducers = self.SCENARIO_REDUCERS if table == "scenarios" else {}
                extras = "".join(f", {name} REAL" for name in reducers)
                self.connection.execute(f"""
                    CREATE TABLE {table} (
                        {definitions}, requests INTEGER NOT NULL, successes INTEGER NOT NULL,
                        total_ms REAL NOT NULL, max_ms REAL NOT NULL, success_ms REAL NOT NULL,
                        first TEXT NOT NULL, last TEXT NOT NULL {extras}, PRIMARY KEY ({names})
                    ) WITHOUT ROWID
                """)
                placeholders = ", ".join("?" for column in range(len(keys) + 7 + len(reducers)))
                updates = ""
                for name, reducer in reducers.items():
                    operation = (f"{name} + excluded.{name}" if reducer == "SUM"
                                 else f"MAX({name}, excluded.{name})")
                    updates += f""", {name} = CASE
                        WHEN {name} IS NULL THEN excluded.{name}
                        WHEN excluded.{name} IS NULL THEN {name}
                        ELSE {operation} END"""
                self.upserts[table] = f"""
                    INSERT INTO {table} VALUES ({placeholders})
                    ON CONFLICT ({names}) DO UPDATE SET
                        requests = requests + excluded.requests,
                        successes = successes + excluded.successes,
                        total_ms = total_ms + excluded.total_ms,
                        max_ms = MAX(max_ms, excluded.max_ms),
                        success_ms = success_ms + excluded.success_ms,
                        first = MIN(first, excluded.first), last = MAX(last, excluded.last)
                        {updates}
                """
            definitions = ", ".join(f"{name} {kind} NOT NULL" for name, kind in self.RECOVERY_KEYS)
            names = ", ".join(name for name, kind in self.RECOVERY_KEYS)
            self.connection.execute(f"""
                CREATE TABLE recovery_episodes (
                    {definitions}, requests INTEGER NOT NULL DEFAULT 0,
                    successes INTEGER NOT NULL DEFAULT 0, cancelled INTEGER NOT NULL DEFAULT 0,
                    events INTEGER NOT NULL DEFAULT 0, first_failure TEXT, first_success TEXT,
                    confirmed TEXT, elapsed REAL, last_event TEXT,
                    outcome TEXT NOT NULL DEFAULT 'onbekend', PRIMARY KEY ({names})
                ) WITHOUT ROWID
            """)
        except BaseException:
            self.close()
            raise

    def add(self, table, key, moment, latency, success, *, request_count=1, extras=()):
        timestamp = moment.isoformat(timespec="microseconds")
        values = tuple(excel_value(value) for value in key)
        self.connection.execute(self.upserts[table], values + (
            request_count, int(success), latency, latency, latency if success else 0,
            timestamp, timestamp,
        ) + tuple(extras))
        self.flush_batch()

    def flush_batch(self):
        self.pending += 1
        if self.pending >= 1000:
            self.connection.commit()
            self.pending = 0

    def rows(self, table):
        self.connection.commit()
        names = ", ".join(name for name, kind in self.KEYS[table])
        for row in self.connection.execute(f"SELECT * FROM {table} ORDER BY {names}"):
            key_count = len(self.KEYS[table])
            metrics = Metrics(*row[key_count:key_count + 5],
                              datetime.fromisoformat(row[key_count + 5]),
                              datetime.fromisoformat(row[key_count + 6]))
            extras = list(row[key_count + 7:])
            if table == "scenarios":
                extras[8] = round(extras[8] / extras[10], 2) if extras[10] else None
                for position in range(11, len(extras), 2):
                    extras[position] = round(extras[position] / extras[position + 1], 2) if extras[position + 1] else None
            values = metrics.values()
            if not metrics.requests:
                values[3:7] = [None] * 4
            yield list(row[:key_count]) + values + extras

    def add_recovery(self, key, row, moment, success):
        names = ", ".join(name for name, kind in self.RECOVERY_KEYS)
        where = " AND ".join(f"{name} = ?" for name, kind in self.RECOVERY_KEYS)
        self.connection.execute(
            f"INSERT OR IGNORE INTO recovery_episodes ({names}) VALUES (?, ?, ?, ?, ?)", key,
        )
        if row["record_type"] != "recovery" and measurement_type(row["phase"]) == "recovery":
            self.connection.execute(f"""
                UPDATE recovery_episodes SET requests = requests + 1,
                    successes = successes + ?, cancelled = cancelled + ? WHERE {where}
            """, (int(success), int(row["record_type"] == "cancelled")) + key)
        elif row["record_type"] == "recovery":
            timestamp = moment.isoformat(timespec="microseconds")
            self.connection.execute(f"""
                UPDATE recovery_episodes SET events = events + 1,
                    first_failure = COALESCE(?, first_failure),
                    first_success = COALESCE(?, first_success), confirmed = COALESCE(?, confirmed),
                    elapsed = CASE WHEN ? >= COALESCE(last_event, '') THEN ? ELSE elapsed END,
                    outcome = CASE WHEN ? >= COALESCE(last_event, '') THEN ? ELSE outcome END,
                    last_event = MAX(COALESCE(last_event, ''), ?) WHERE {where}
            """, tuple(row[column].isoformat(timespec="microseconds") if row[column] else None for column in
                       ("first_failure_utc", "first_success_utc", "confirmed_utc")) + (
                timestamp, row["recovery_elapsed_s"], timestamp, row["recovery_outcome"], timestamp,
            ) + key)
        self.flush_batch()

    def recovery_rows(self):
        self.connection.commit()
        query = """
            WITH episodes AS (
                SELECT *, COALESCE(elapsed,
                    CASE WHEN first_failure IS NOT NULL AND last_event IS NOT NULL
                            AND julianday(COALESCE(confirmed, last_event)) >= julianday(first_failure)
                        THEN (julianday(COALESCE(confirmed, last_event)) - julianday(first_failure)) * 86400
                    END) AS duration FROM recovery_episodes
            )
            SELECT run_id, host, endpoint, scenario, 'recovery', outcome,
                COUNT(*), SUM(events), SUM(requests), SUM(successes),
                SUM(requests - successes), SUM(cancelled), COUNT(duration),
                ROUND(AVG(duration), 2), ROUND(MIN(duration), 2), ROUND(MAX(duration), 2),
                MIN(first_failure), MAX(last_event)
            FROM episodes GROUP BY run_id, host, endpoint, scenario, outcome
            ORDER BY run_id, host, endpoint, scenario, outcome
        """
        for row in self.connection.execute(query):
            yield list(row[:-2]) + [datetime.fromisoformat(value) if value else None for value in row[-2:]]

    def close(self):
        try:
            if self.connection is not None:
                self.connection.close()
        finally:
            self.directory.cleanup()


def snapshot_lines(path):
    with path.open("rb") as source:
        # Bound the snapshot so a growing log cannot keep the export reading forever.
        remaining = os.fstat(source.fileno()).st_size
        first_line = True
        while remaining:
            line = source.readline(remaining)
            if not line:
                break
            remaining -= len(line)
            if not line.endswith(b"\n"):
                break
            yield line.decode("utf-8-sig" if first_line else "utf-8")
            first_line = False


def excel_value(value):
    if isinstance(value, datetime):
        return value.isoformat(timespec="milliseconds") + "Z"
    return value


def utc_moment(value):
    moment = datetime.fromisoformat(value)
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(timezone.utc).replace(tzinfo=None)


def optional_number(value, column, *, integer=False):
    if value in (None, ""):
        return None
    try:
        number = int(value) if integer else float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {column}: expected a nonnegative {'integer' if integer else 'number'}") from exc
    if number < 0 or (not integer and not math.isfinite(number)):
        raise ValueError(f"Invalid {column}: expected a finite nonnegative number")
    if integer and number > 9_223_372_036_854_775_807:
        raise ValueError(f"Invalid {column}: outside SQLite integer range")
    return number


def parse_record(row):
    if None in row or any(value is None for value in row.values()):
        raise ValueError("Incomplete or extra CSV fields")
    parsed = dict(row)
    parsed["timestamp"] = utc_moment(row["timestamp"])
    parsed["record_type"] = row.get("record_type") or "request"
    if parsed["record_type"] not in {"request", "cancelled", "recovery"}:
        raise ValueError("Invalid record_type")
    event_record = parsed["record_type"] == "recovery"
    parsed["latency_ms"] = optional_number(row["latency_ms"], "latency")
    for column in ("request", "http_status"):
        parsed[column] = int(row[column]) if row[column] else None
    if not event_record and any(parsed[column] is None for column in
                                ("request", "http_status", "latency_ms")):
        raise ValueError("Missing request numeric values")
    parsed["repeat"] = int(row["repeat"]) if row.get("repeat") else 1
    for column in FLOAT_DIAGNOSTICS | INTEGER_DIAGNOSTICS:
        parsed[column] = optional_number(row.get(column), column, integer=column in INTEGER_DIAGNOSTICS)
    for column in RECOVERY_TIMESTAMPS:
        parsed[column] = utc_moment(row[column]) if row.get(column) else None
    if event_record:
        event = row.get("recovery_event")
        if not row.get("recovery_id") or event not in RECOVERY_EVENTS:
            raise ValueError("Recovery record requires recovery_id and a valid recovery_event")
        parsed["recovery_outcome"] = row.get("recovery_outcome") or (
            event if event in {"recovered", "timeout", "interrupted"} else "in_progress"
        )
    return parsed


def measurement_type(phase):
    prefix = phase.split("_", 1)[0]
    return prefix if prefix in {"health", "recovery"} else "stress"


def connection_type(row):
    new = row["new_connections"]
    established = row["established_connections"]
    reused = row["reused_connections"]
    has_new = bool(new or established)
    if has_new and reused:
        return "mixed"
    if has_new and reused == 0:
        return "new"
    if reused and new == 0 and established in (None, 0):
        return "reused"
    return "unknown"


class SheetWriter:
    def __init__(self, workbook, title, headers, max_rows, widths=None):
        from openpyxl.styles import Alignment, Font, PatternFill

        if not 2 <= max_rows <= EXCEL_MAX_ROWS:
            raise ValueError("Sheet row limit must be between 2 and 1048576")
        self.workbook = workbook
        self.title = title
        self.headers = headers
        self.max_rows = max_rows
        self.widths = widths or {}
        self.text_alignment = Alignment(vertical="center", wrap_text=True)
        self.number_alignment = Alignment(horizontal="right", vertical="center")
        self.band_fill = PatternFill("solid", fgColor="F1F5F4")
        self.body_font = Font(name="Calibri", size=11, color="24332D")
        self.alert_font = Font(name="Calibri", size=11, color="B42318", bold=True)
        self.success_font = Font(name="Calibri", size=11, color="217346", bold=True)
        self.sheets = []
        self.row_count = 0
        self.new_sheet()

    def new_sheet(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        if self.sheets:
            self.finish_sheet()
        number = len(self.sheets) + 1
        name = self.title if number == 1 else f"{self.title}_{number}"
        self.sheet = self.workbook.create_sheet(name)
        self.sheets.append(self.sheet)
        self.sheet.freeze_panes = "A2"
        self.sheet.sheet_view.showGridLines = False
        self.sheet.sheet_view.zoomScale = 85
        self.sheet.sheet_properties.pageSetUpPr.fitToPage = True
        self.sheet.page_setup.orientation = "landscape"
        self.sheet.page_setup.fitToWidth = 1
        self.sheet.page_setup.fitToHeight = 0
        self.sheet.print_title_rows = "1:1"
        self.sheet.sheet_format.defaultRowHeight = 32
        self.sheet.row_dimensions[1].height = 44
        self.sheet.sheet_properties.tabColor = (
            "B42318" if self.title in {"Fouten", "Foutdiagnose"} else
            "B7791F" if "herstel" in self.title.lower() else "256D5B"
        )
        for column, header in enumerate(self.headers, 1):
            width = self.widths.get(header, 28 if "UTC" in header else min(38, max(16, len(header) + 3)))
            self.sheet.column_dimensions[get_column_letter(column)].width = width
        cells = []
        for header in self.headers:
            cell = WriteOnlyCell(self.sheet, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="256D5B")
            cell.alignment = self.text_alignment
            cells.append(cell)
        self.sheet.append(cells)
        self.row_count = 1

    def append(self, values, *, formats=None):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

        if len(values) != len(self.headers):
            raise ValueError("Worksheet row does not match its header")
        if self.row_count >= self.max_rows:
            self.new_sheet()
        cells = []
        for header, value in zip(self.headers, values):
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell = WriteOnlyCell(self.sheet, value=excel_value(value))
            cell.font = self.body_font
            cell.alignment = self.number_alignment if isinstance(value, (int, float)) else self.text_alignment
            if self.row_count % 2:
                cell.fill = self.band_fill
            if header in {"Fouten", "Geannuleerd"} and isinstance(value, (int, float)) and value > 0:
                cell.font = self.alert_font
            if header == "Hersteluitkomst":
                if value in {"temporary", "transient", "recovered"}:
                    cell.font = self.success_font
                elif value in {"not_recovered", "timeout", "interrupted"}:
                    cell.font = self.alert_font
            if isinstance(value, str) or isinstance(value, datetime):
                # Log text must remain literal data, even when it starts with '='.
                cell.data_type = "s"
            if header == "Foutpercentage":
                cell.number_format = "0.00%"
            elif header.endswith(" ms") or header in {
                "Connectie ms (DNS+TCP+TLS)", "Gem. connectie ms (DNS+TCP+TLS)", "Herstel verstreken s",
                "Gem. herstel s", "Min. herstel s", "Max. herstel s",
                "Gem. actieve requests bij start",
            }:
                cell.number_format = "0.00"
            if formats and header in formats:
                cell.number_format = formats[header]
            cells.append(cell)
        self.sheet.append(cells)
        self.row_count += 1

    def finish_sheet(self):
        from openpyxl.utils import get_column_letter

        self.sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(self.headers))}{self.row_count}"
        )


def add_trend_charts(writer):
    from openpyxl.chart import LineChart, Reference

    sheet = writer.sheets[0]
    if writer.row_count < 3:
        return
    for column, title, anchor in [
        ("Gem. ms", "Gemiddelde responstijd per uur", "L2"),
        ("Foutpercentage", "Foutpercentage per uur", "L18"),
    ]:
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = "ms" if column == "Gem. ms" else "Fractie fouten"
        chart.x_axis.title = "Uur UTC"
        chart.add_data(Reference(
            sheet, min_col=writer.headers.index(column) + 1,
            min_row=1, max_row=writer.row_count,
        ), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=writer.row_count))
        chart.width = 23
        chart.height = 8
        sheet.add_chart(chart, anchor)


def export_report(inputs, output, *, max_rows=EXCEL_MAX_ROWS):
    from openpyxl import Workbook

    paths = [Path(source).resolve() for source in inputs]
    if not paths or len(set(paths)) != len(paths):
        raise ValueError("Provide at least one CSV, without duplicate input files")
    output = Path(output).resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Excel output must have the .xlsx extension")
    if output in paths:
        raise ValueError("Excel output cannot overwrite an input CSV")
    for path in paths:
        with path.open("r", newline="", encoding="utf-8-sig") as source:
            fields = next(csv.reader(source), [])
        if not REQUIRED_COLUMNS.issubset(fields):
            raise ValueError(f"Missing required CSV columns in {path}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    overview = SheetWriter(workbook, "Overzicht", ["Kenmerk", "Waarde"], max_rows,
                           widths={"Kenmerk": 30, "Waarde": 95})
    errors = SheetWriter(workbook, "Fouten", ERROR_COLUMNS, max_rows,
                         widths={"Fout": 90, "Bronbestand": 60, "Run ID": 38})
    recovery = SheetWriter(workbook, "Herstel", RECOVERY_COLUMNS, max_rows,
                           widths={"Run ID": 38, "Recovery ID": 38, "Bronbestand": 60})
    totals = Metrics()
    aggregates = None
    temp_path = None
    try:
        aggregates = DiskAggregates()
        for path in paths:
            reader = csv.DictReader(snapshot_lines(path), strict=True)
            for row in reader:
                try:
                    if None in row or any(row.get(column) is None for column in REQUIRED_COLUMNS):
                        raise ValueError("Incomplete or extra CSV fields")
                    row = parse_record(row)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"{path}, CSV line {reader.line_num}: {exc}") from exc
                moment = row["timestamp"]
                latency = row["latency_ms"]
                status = row["http_status"]
                request = row["request"]
                repeat = row["repeat"]
                record_type = row["record_type"]
                is_request = record_type != "recovery"
                cancelled = record_type == "cancelled"
                success = is_request and not cancelled and status == 200 and not row["error"]
                host = row.get("host") or "onbekend"
                endpoint = row.get("endpoint") or "onbekend"
                run_id = row.get("run_id") or f"legacy:{path}"
                phase = row["phase"]
                protocol = row["protocol"]
                scenario = row.get("scenario") or (
                    "onbekend" if measurement_type(phase) != "stress" else phase
                )
                row["scenario"] = scenario
                scenario_key = (run_id, host, endpoint, scenario, measurement_type(phase), record_type,
                                protocol, str(row["batch_size"]) if row["batch_size"] is not None else "onbekend",
                                str(row["interval_s"]) if row["interval_s"] is not None else "onbekend",
                                row.get("capture_manifest") or "onbekend")
                active = row["active_requests"] if is_request else None
                counts = [row[column] if is_request else None for column in
                          ("new_connections", "established_connections", "reused_connections")]
                aggregates.add("scenarios", scenario_key, moment, latency if is_request else 0,
                               success, request_count=int(is_request), extras=(
                    *counts, *(int(value is not None) for value in counts),
                    int(cancelled), int(not is_request), active, active, int(active is not None),
                    *[value for column in STAGE_COLUMNS for value in (
                        row[column] if is_request else None, int(is_request and row[column] is not None),
                    )],
                ))
                if row.get("recovery_id"):
                    aggregates.add_recovery((run_id, host, endpoint, scenario, row["recovery_id"]),
                                            row, moment, success)
                if not is_request:
                    recovery.append([
                        moment, run_id, repeat, host, endpoint, row["recovery_id"], scenario,
                        row["recovery_event"], row["first_failure_utc"], row["first_success_utc"],
                        row["confirmed_utc"], row["recovery_elapsed_s"], row["recovery_success_streak"],
                        row["recovery_failure_streak"], row["recovery_outcome"], str(path),
                    ])
                    continue
                hour = moment.replace(minute=0, second=0, microsecond=0)
                day = moment.replace(hour=0, minute=0, second=0, microsecond=0)
                totals.add(moment, latency, success)
                for table, key in [
                    ("hours", (hour,)),
                    ("hourly_phases", (hour, host, endpoint, phase, protocol)),
                    ("daily_phases", (day, host, endpoint, phase, protocol)),
                    ("repeats", (run_id, repeat, host, endpoint)),
                ]:
                    aggregates.add(table, key, moment, latency, success)
                if not success:
                    connection = connection_type(row)
                    aggregates.add("diagnostics", scenario_key + (
                        row.get("error_stage") or "onbekend",
                        row.get("exception_type") or "onbekend", connection,
                        str(active) if active is not None else "onbekend",
                    ), moment, latency, False)
                    errors.append([
                        moment, run_id, repeat, host, endpoint, phase, protocol,
                        request, status, latency, row["error"], str(path),
                    ] + [row.get(column) if row.get(column) != "" else None
                         for column in DIAGNOSTIC_COLUMNS] + [connection])

        for label, value in [
            ("Gegenereerd UTC", datetime.now(timezone.utc).replace(tzinfo=None)),
            ("Bronbestanden", len(paths)),
            ("Requests", totals.requests),
            ("OK", totals.successes),
            ("Fouten", totals.requests - totals.successes),
            ("Foutpercentage", totals.values()[3]),
            ("Gemiddelde ms (alle)", totals.values()[4]),
            ("Maximale ms (alle)", totals.values()[5]),
            ("Gemiddelde ms (OK)", totals.values()[6]),
            ("Eerste meting UTC", totals.first),
            ("Laatste meting UTC", totals.last),
            ("Tijdzone", "UTC"),
            ("Detailmetingen", "Alle requests blijven beschikbaar in de bron-CSV."),
            ("Trend", "Uurtrend combineert alle fases, protocollen en hosts; filter Uren voor details."),
            ("Foutdefinitie", "HTTP-status niet 200, of een fout tijdens het uitlezen van de response."),
            ("Geannuleerd", "Geannuleerde requests tellen als fout; Scenarios vermeldt aparte aantallen."),
            ("Herstel", "Herstelevents tellen niet als requests; Herstel toont events in bronvolgorde."),
            ("Connectietiming", "Connectie ms omvat DNS + TCP + TLS; DNS ms is hiervan een deelmeting."),
            ("Pakketregistratie", "Scenarios vermeldt capturemetadata; PCAP toont TCP en zichtbare TLS-berichten, zonder decryptie."),
            ("Diagnostiek", "Lege diagnostiek is onbekend; meetkolommen tonen de dekking van deelsommen."),
            ("Actieve requests", "Aantal actieve requests bij requeststart, niet de ingestelde concurrency."),
            ("Herstelduur", "Scenarioherstel groepeert per herstel-ID en uitkomst; in_progress is nog niet afgerond."),
        ]:
            overview.append([label, value], formats={"Waarde": "0.00%"} if label == "Foutpercentage" else None)
        for path in paths:
            overview.append(["Bronbestand", str(path)])

        writers = [overview, errors, recovery]
        for title, headers, table in [
            ("Uurtrend", ["Uur UTC"], "hours"),
            ("Uren", ["Uur UTC", "Host", "Endpoint", "Fase", "Protocol"], "hourly_phases"),
            ("Dagen", ["Dag UTC", "Host", "Endpoint", "Fase", "Protocol"], "daily_phases"),
            ("Herhalingen", ["Run ID", "Repeat", "Host", "Endpoint"], "repeats"),
            ("Scenarios", SCENARIO_HEADERS, "scenarios"),
            ("Foutdiagnose", SCENARIO_HEADERS + [
                "Foutfase", "Exceptiontype", "Connectietype", "Actieve requests bij start",
            ], "diagnostics"),
        ]:
            extras = SCENARIO_EXTRA_HEADERS if table == "scenarios" else []
            writer = SheetWriter(workbook, title, headers + METRIC_COLUMNS + extras, max_rows,
                                 widths={"Run ID": 38})
            for values in aggregates.rows(table):
                writer.append(values)
            if title == "Uurtrend" and len(writer.sheets) == 1:
                add_trend_charts(writer)
            writers.append(writer)
        recovery_stats = SheetWriter(workbook, "Scenarioherstel", [
            "Run ID", "Host", "Endpoint", "Scenario", "Recordtype", "Hersteluitkomst",
            "Herstelseries", "Herstelevents", "Herstelrequests", "OK", "Fouten", "Geannuleerd",
            "Duurmetingen", "Gem. herstel s", "Min. herstel s", "Max. herstel s",
            "Eerste fout UTC", "Laatste event UTC",
        ], max_rows, widths={"Run ID": 38})
        for values in aggregates.recovery_rows():
            recovery_stats.append(values)
        writers.append(recovery_stats)
        for writer in writers:
            writer.finish_sheet()

        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temporary:
            temp_path = Path(temporary.name)
        workbook.save(temp_path)
        os.replace(temp_path, output)
        return totals.requests
    finally:
        try:
            for sheet in workbook.worksheets:
                if not sheet.closed:
                    sheet.close()
                writer = sheet._writer
                if writer and os.path.exists(writer.out):
                    writer.cleanup()
            workbook.close()
        finally:
            try:
                if aggregates is not None:
                    aggregates.close()
            finally:
                if temp_path and temp_path.exists():
                    temp_path.unlink()


def expand_inputs(patterns):
    paths = []
    for pattern in patterns:
        matches = sorted(glob.glob(pattern)) if glob.has_magic(pattern) else [pattern]
        if not matches:
            raise ValueError(f"No CSV files match {pattern}")
        paths.extend(matches)
    return paths


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="*", default=[str(DEFAULT_INPUT)], help="CSV paths or quoted glob patterns")
    parser.add_argument("--output", help="Excel path; default is the first CSV with .xlsx extension")
    args = parser.parse_args()
    try:
        paths = expand_inputs(args.inputs)
        output = Path(args.output) if args.output else Path(paths[0]).with_suffix(".xlsx")
        requests = export_report(paths, output)
    except ImportError as exc:
        parser.exit(1, f"Report failed: {exc}\nInstall dependencies with: python -m pip install -r requirements.txt\n")
    except (OSError, ValueError, csv.Error, sqlite3.Error) as exc:
        parser.exit(1, f"Report failed: {exc}\n")
    print(f"Excel: {output} ({requests} requests)")


if __name__ == "__main__":
    main()